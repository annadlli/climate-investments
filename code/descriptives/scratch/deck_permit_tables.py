"""
Authors: Claude (for Anna Li)
Date: 2026-10-14
Claude change 10-14 (new file; Anna: three tables for a new deck workbook: (1) Builty vs ATTOM
elevation permits per state on NFIP-matched SFHA homes with percentages, (2) a trace of the ATTOM
permits through every match tier, (3) permit status in the final dataset: total, done, live,
dead, plus a column that allows for feeds whose last-known label never reaches a completion).

Scratch script, run by hand, not in master.do. Writes a NEW workbook
    Dropbox Flooding/Decks/Deck_Oct 14.xlsx
with three tabs. The Oct 2 deck is not copied because it carries embedded figures that an
openpyxl round-trip drops; drag the tabs across in Excel. Deck conventions followed: Calibri 12,
bold title in A1, one-line note in A2, bold header row, counts #,##0, shares 0.00% as formulas.

Inputs (all on Dropbox unless noted):
  build/nfip_attom_pipeline_v2/nfip_attom_property/{st}_nfip_attom_property.parquet   10-12 ATTOM links
  build/nfip_attom_pipeline_v2/nfip_attom_property/links_0919/{st}_...parquet           09-19 Builty links
  build/nfip_attom_pipeline_v2/attom_permits/{st}_permits_matched.parquet               permit -> ATTOM tiers
  clean/nfip_policies_state/{st}.dta                                                    sfha by property
  clean/permit_elevations.dta                                                           permit addresses, status
  <scratch>/raw_status_by_jurisdiction.parquet  (optional; from the raw permit scan: state, county,
      city, status, n over every raw permit in FL LA TX) for the feed-closes column of table 3.
Usage: python deck_permit_tables.py --data <Data root> --out <xlsx path> [--scan <parquet>]
"""
import argparse, re, sys
from pathlib import Path
import duckdb, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).parent))
from permit_status_done import bucket   # the 10-14 done / live / dead rule

STATES = ["FL", "LA", "TX"]
NAMES = {"FL": "Florida", "LA": "Louisiana", "TX": "Texas"}
FONT = Font(name="Calibri", size=12); BOLD = Font(name="Calibri", size=12, bold=True)
CNT = "#,##0"; PCT = "0.00%"
DONE_RE = re.compile(r"\b(final|finaled|finalled|complete|completed|closed)\b")   # feed-closes test
FEED_FLOOR = 0.05   # a feed "records completions" when >= 5% of all its permits ever carry a done label
# Claude change 10-14 (later): the grain is the city feed (state, county, city), since ATTOM sources permits by
# jurisdiction; at the county grain almost every county passes a 1% floor while whole city feeds (Morgan City,
# Houma, Chalmette ...) never write a completion. 5% separates the two: done addresses sit in feeds with a done
# share of 0.15 at the 5th percentile, live addresses in LA mostly in feeds below 0.05.
def norm(x):
    return re.sub(r"[^A-Z0-9 ]", "", str(x).upper()).strip()

def put(ws, r, c, v, bold=False, fmt=None):
    cell = ws.cell(row=r, column=c, value=v); cell.font = BOLD if bold else FONT
    if fmt: cell.number_format = fmt
    return cell

def header(ws, row, labels):
    for j, lab in enumerate(labels, start=1):
        put(ws, row, j, lab, bold=True).alignment = Alignment(wrap_text=True, vertical="top")

def widths(ws, first=14, rest=16):
    ws.column_dimensions["A"].width = first
    for j in range(2, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(j)].width = rest

# ----------------------------------------------------------------------------- data
def sfha_by_property(data, st):
    pol = pd.read_stata(f"{data}/clean/nfip_policies_state/{st.lower()}.dta", columns=["property_id", "sfha"])
    return pol.groupby("property_id").sfha.max().reset_index().rename(columns={"property_id": "property_id_state"})

def table1(con, data):
    P = f"{data}/build/nfip_attom_pipeline_v2/nfip_attom_property"
    rows = []
    for st in STATES:
        sf = sfha_by_property(data, st)
        for src, path, flag, retro in [("Builty", f"{P}/links_0919", "builty_elevated", "builty_retrofit"),
                                       ("ATTOM", P, "permit_elevated", "permit_retrofit")]:
            link = con.execute(f"""select property_id_state, coalesce({flag},0)=1 as elev,
                coalesce({retro},0)=1 as retro, match_tier_number is not null as matched
                from '{path}/{st.lower()}_nfip_attom_property.parquet'""").df()
            m = link.merge(sf, on="property_id_state", how="left"); s = m[(m.sfha == 1) & m.matched]
            rows.append(dict(state=st, source=src, nfip=len(m), sfha=int((m.sfha == 1).sum()),
                             sfha_matched=len(s), permit=int(s.elev.sum()), retrofit=int(s.retro.sum())))
    return pd.DataFrame(rows)

def table2(con, data, pe):
    P = f"{data}/build/nfip_attom_pipeline_v2"
    a, b = [], []
    for st in STATES:
        pm = con.execute(f"select attom_match_tier from '{P}/attom_permits/{st.lower()}_permits_matched.parquet'").df()
        t = pm.attom_match_tier.value_counts().to_dict()
        a.append(dict(state=st, addresses=int((pe.state == st).sum()), retrofit=int(((pe.state == st) & (pe.retrofit == 1)).sum()),
                      with_attomid=int((pe.state == st).sum() - (pe[pe.state == st].attomid == "").sum()), **t))
        link = con.execute(f"""select match_tier_number as tier, match_tier as name, count(*) as n
            from '{P}/nfip_attom_property/{st.lower()}_nfip_attom_property.parquet'
            where permit_elevated = 1 group by 1, 2""").df()
        link["state"] = st; b.append(link)
    return pd.DataFrame(a), pd.concat(b)

def feed_closes(scan):
    """(state, county, city) -> closes = True when the city feed records completion labels on >= FEED_FLOOR of
    all its raw permits (every permit type). Keys are normalized (upper case, letters and digits only) and the
    raw rows aggregated AFTER normalizing, so spelling variants of one feed do not duplicate the key."""
    if scan is None or not Path(scan).exists(): return None
    d = pd.read_parquet(scan)
    d["county"] = d.county.map(norm); d["city"] = d.city.map(norm)
    d["done_n"] = d.n * d.status.fillna("").map(lambda s: bool(DONE_RE.search(s)))
    g = d.groupby(["state", "county", "city"], as_index=False).agg(n=("n", "sum"), done_n=("done_n", "sum"))
    g["share_done"] = g.done_n / g.n
    g["closes"] = g.share_done >= FEED_FLOOR
    return g

def table3(con, data, pe, fc):
    P = f"{data}/build/nfip_attom_pipeline_v2"
    rows = []
    def add(st, level, b, closes):
        done = int((b == "done").sum()); live = int((b == "live").sum()); dead = int((b == "dead").sum())
        adj = None if closes is None else done + int(((b == "live") & ~closes).sum())
        rows.append(dict(state=st, level=level, total=len(b), done=done, live=live, dead=dead, done_adj=adj))
    for st in STATES:
        p = pe[pe.state == st].copy(); p["b"] = p.status.map(bucket)
        cl = None
        if fc is not None:   # Claude change 10-14 (later): city-feed grain
            key = fc[fc.state == st].set_index(["county", "city"]).closes
            cl = pd.Series(list(zip(p.county.map(norm), p.locality.map(norm))), index=p.index).map(key).fillna(True)   # unknown feed: closing
        add(st, "Elevation permit addresses (permit file)", p.b, cl)
        pm = con.execute(f"""select ATTOMID, COUNTY, LOCALITY, STATUS, attom_match_tier from '{P}/attom_permits/{st.lower()}_permits_matched.parquet'
                             where attom_match_tier <> 'unmatched'""").df()
        pm["b"] = pm.STATUS.map(bucket)
        clm = None
        if fc is not None:   # Claude change 10-14 (later): city-feed grain
            key = fc[fc.state == st].set_index(["county", "city"]).closes
            clm = pd.Series(list(zip(pm.COUNTY.map(norm), pm.LOCALITY.map(norm))), index=pm.index).map(key).fillna(True)
        add(st, "Permit addresses linked to an ATTOM property", pm.b, clm)
        link = con.execute(f"""select property_id_state, cast(assigned_attomid as varchar) as attomid
            from '{P}/nfip_attom_property/{st.lower()}_nfip_attom_property.parquet' where permit_elevated = 1""").df()
        pm["ATTOMID"] = pm.ATTOMID.astype(str)
        rank = {"done": 0, "live": 1, "dead": 2}   # an address with several permits takes its best status
        best = pm.assign(r=pm.b.map(rank)).sort_values("r").drop_duplicates("ATTOMID").set_index("ATTOMID")
        link["b"] = link.attomid.map(best.b)
        lc = None if clm is None else link.attomid.map(clm.groupby(pm.ATTOMID).max()).fillna(True)
        add(st, "NFIP homes with a permit", link.b, lc)
        sf = sfha_by_property(data, st); link = link.merge(sf, on="property_id_state", how="left")
        ls = link[link.sfha == 1]
        add(st, "NFIP homes with a permit, SFHA", ls.b, None if lc is None else lc[link.sfha == 1])
    return pd.DataFrame(rows)

# ----------------------------------------------------------------------------- sheets
def sheet1(wb, t1):
    ws = wb.create_sheet("Builty vs ATTOM SFHA")
    put(ws, 1, 1, "Home elevation permits by source, NFIP homes in an SFHA that are matched to an ATTOM property", bold=True)
    put(ws, 2, 1, "Builty = 09-19 link files (last Builty match); ATTOM = 10-12 link files (zone-aligned match). SFHA = the home's own "
                  "NFIP rating, ever flagged. Denominator = SFHA homes with an ATTOM match under each source's links, since a permit "
                  "can only be observed on a matched home. Shares are percentages of that denominator. Retrofit excludes elevated new construction.")
    header(ws, 4, ["State", "NFIP homes", "SFHA homes", "Builty: SFHA matched", "Builty permit homes", "Builty share",
                   "ATTOM: SFHA matched", "ATTOM permit homes", "ATTOM share", "ATTOM / Builty",
                   "Builty retrofit homes", "Builty retrofit share", "ATTOM retrofit homes", "ATTOM retrofit share"])
    r = 5
    for st in STATES:
        b = t1[(t1.state == st) & (t1.source == "Builty")].iloc[0]; a = t1[(t1.state == st) & (t1.source == "ATTOM")].iloc[0]
        put(ws, r, 1, NAMES[st]); put(ws, r, 2, int(b.nfip), fmt=CNT); put(ws, r, 3, int(b.sfha), fmt=CNT)
        put(ws, r, 4, int(b.sfha_matched), fmt=CNT); put(ws, r, 5, int(b.permit), fmt=CNT); put(ws, r, 6, f"=E{r}/D{r}", fmt=PCT)
        put(ws, r, 7, int(a.sfha_matched), fmt=CNT); put(ws, r, 8, int(a.permit), fmt=CNT); put(ws, r, 9, f"=H{r}/G{r}", fmt=PCT)
        put(ws, r, 10, f"=I{r}/F{r}", fmt="0.0")
        put(ws, r, 11, int(b.retrofit), fmt=CNT); put(ws, r, 12, f"=K{r}/D{r}", fmt=PCT)
        put(ws, r, 13, int(a.retrofit), fmt=CNT); put(ws, r, 14, f"=M{r}/G{r}", fmt=PCT)
        r += 1
    put(ws, r, 1, "All three", bold=True)
    for c in [2, 3, 4, 5, 7, 8, 11, 13]:
        L = get_column_letter(c); put(ws, r, c, f"=SUM({L}5:{L}{r-1})", bold=True, fmt=CNT)
    put(ws, r, 6, f"=E{r}/D{r}", bold=True, fmt=PCT); put(ws, r, 9, f"=H{r}/G{r}", bold=True, fmt=PCT)
    put(ws, r, 10, f"=I{r}/F{r}", bold=True, fmt="0.0"); put(ws, r, 12, f"=K{r}/D{r}", bold=True, fmt=PCT); put(ws, r, 14, f"=M{r}/G{r}", bold=True, fmt=PCT)
    widths(ws); ws.row_dimensions[4].height = 48

def sheet2(wb, a, b):
    ws = wb.create_sheet("ATTOM permit trace")
    put(ws, 1, 1, "Trace of the ATTOM elevation permits: permit address -> ATTOM property (stage A) -> NFIP home (stage B)", bold=True)
    put(ws, 2, 1, "Stage A: clean/permit_elevations.dta addresses (FL LA TX) and the tier on which build/attom_permits.py linked each to an ATTOM "
                  "property: the ATTOMID ATTOM itself assigned, then the exact-address tiers, then the loose house-number key and Jaro-Winkler. "
                  "Stage B: NFIP homes carrying a permit in the 10-12 link files, by the Wagner cell tier that assigned the ATTOM house "
                  "(bg = block group; zone = flood zone; risk = high/low risk; exact year / 5-year band of construction year; tiers 11-14 have no flood zone; 15 = no construction year).")
    tiersA = ["attomid", "exact", "unique_no_suffix", "unique_addr_county", "unique_compact", "unique_nosuffix_county", "loose_key", "jaro_winkler", "unmatched"]
    labelsA = {"attomid": "ATTOMID on the permit", "exact": "Exact address", "unique_no_suffix": "Unique, no suffix", "unique_addr_county": "Unique address in county",
               "unique_compact": "Unique compact address", "unique_nosuffix_county": "Unique no-suffix in county", "loose_key": "Loose house-number key",
               "jaro_winkler": "Jaro-Winkler", "unmatched": "Unmatched"}
    put(ws, 4, 1, "Stage A: permit addresses to ATTOM properties", bold=True)
    header(ws, 5, ["Tier"] + [NAMES[s] for s in STATES] + ["All three", "Share"])
    r = 6
    put(ws, r, 1, "Elevation permit addresses"); [put(ws, r, 2 + i, int(a[a.state == s].addresses.iloc[0]), fmt=CNT) for i, s in enumerate(STATES)]
    put(ws, r, 5, f"=SUM(B{r}:D{r})", fmt=CNT); put(ws, r, 6, f"=E{r}/E$6", fmt=PCT); r += 1
    put(ws, r, 1, "of which retrofit addresses"); [put(ws, r, 2 + i, int(a[a.state == s].retrofit.iloc[0]), fmt=CNT) for i, s in enumerate(STATES)]
    put(ws, r, 5, f"=SUM(B{r}:D{r})", fmt=CNT); put(ws, r, 6, f"=E{r}/E$6", fmt=PCT); r += 1
    put(ws, r, 1, "of which carry an ATTOMID"); [put(ws, r, 2 + i, int(a[a.state == s].with_attomid.iloc[0]), fmt=CNT) for i, s in enumerate(STATES)]
    put(ws, r, 5, f"=SUM(B{r}:D{r})", fmt=CNT); put(ws, r, 6, f"=E{r}/E$6", fmt=PCT); r += 1
    first = r
    for t in tiersA:
        put(ws, r, 1, "  " + labelsA[t])
        for i, s in enumerate(STATES):
            v = a[a.state == s][t].iloc[0] if t in a.columns else 0
            put(ws, r, 2 + i, int(0 if pd.isna(v) else v), fmt=CNT)
        put(ws, r, 5, f"=SUM(B{r}:D{r})", fmt=CNT); put(ws, r, 6, f"=E{r}/E$6", fmt=PCT); r += 1
    put(ws, r, 1, "Linked to an ATTOM property (all tiers but unmatched)", bold=True)
    for i in range(3):
        L = get_column_letter(2 + i); put(ws, r, 2 + i, f"=SUM({L}{first}:{L}{r-2})", bold=True, fmt=CNT)
    put(ws, r, 5, f"=SUM(B{r}:D{r})", bold=True, fmt=CNT); put(ws, r, 6, f"=E{r}/E$6", bold=True, fmt=PCT)
    r += 1; put(ws, r, 1, "Link rate by state"); [put(ws, r, 2 + i, f"={get_column_letter(2+i)}{r-1}/{get_column_letter(2+i)}$6", fmt=PCT) for i in range(3)]
    r += 2
    put(ws, r, 1, "Stage B: NFIP homes carrying a permit, by Wagner cell tier (10-12 links)", bold=True); r += 1
    header(ws, r, ["Tier"] + [NAMES[s] for s in STATES] + ["All three", "Share"]); r += 1
    first = r
    names = b.drop_duplicates("tier").set_index("tier").name.to_dict()
    for t in range(1, 16):
        if t not in names: continue
        put(ws, r, 1, names[t])
        for i, s in enumerate(STATES):
            v = b[(b.state == s) & (b.tier == t)].n
            put(ws, r, 2 + i, int(v.iloc[0]) if len(v) else 0, fmt=CNT)
        put(ws, r, 5, f"=SUM(B{r}:D{r})", fmt=CNT); put(ws, r, 6, f"=E{r}/E${first + 15 - (15 - len(names))}", fmt=PCT); r += 1
    last = r - 1
    put(ws, r, 1, "NFIP homes with a permit", bold=True)
    for i in range(4):
        L = get_column_letter(2 + i); put(ws, r, 2 + i, f"=SUM({L}{first}:{L}{last})", bold=True, fmt=CNT)
    put(ws, r, 6, 1.0, bold=True, fmt=PCT)
    # fix the share formulas now that the total row is known
    for rr in range(first, last + 1):
        ws.cell(row=rr, column=6).value = f"=E{rr}/E${r}"
    r += 1; put(ws, r, 1, "  of which block-group tiers (1-4, 11, 15)")
    bg_rows = [rr for rr in range(first, last + 1) if ws.cell(row=rr, column=1).value.split("_")[0] in ("1", "2", "3", "4", "11", "15")]
    for i in range(4):
        L = get_column_letter(2 + i); put(ws, r, 2 + i, "=" + "+".join(f"{L}{rr}" for rr in bg_rows), fmt=CNT)
    put(ws, r, 6, f"=E{r}/E{r-1}", fmt=PCT)
    widths(ws, first=44)

def sheet3(wb, t3, fc):
    ws = wb.create_sheet("Permit status")
    put(ws, 1, 1, "Status of the ATTOM elevation permits in the final permit dataset: done, live, dead", bold=True)
    note = ("Status = the raw last-known label on the permit record (ATTOM's STATUS; ~18,500 spellings nationally). Done = final / complete / finaled / "
            "closed / CO or COC issued / inspections completed / passed; live = issued / approved / applied / filed / open / in review / pending / "
            "hold / suspended / reinstated; dead = expired / cancelled / withdrawn / rejected / revoked / void / inactive. Classification in "
            "descriptives/scratch/permit_status_done.py. An address with several permits takes its best status.")
    if fc is not None:
        note += (f" 'Done, feed-adjusted' adds live permits in jurisdictions whose feed never records completions: a city feed (state, county, city) "
                 f"whose raw ATTOM permits of every type carry a done label on fewer than {FEED_FLOOR:.0%} of permits, so an 'issued' there is the last "
                 f"label the feed ever writes and says nothing about completion. Feeds not in the scan are treated as closing. Scan: raw "
                 f"attom_permits_all.parquet, FL LA TX, status by state x county x city (Claude change 10-14).")
    put(ws, 2, 1, note)
    cols = ["State", "Level", "Total", "Done", "Live", "Dead", "Done share", "Done, feed-adjusted", "Adjusted share"]
    header(ws, 4, cols); r = 5
    for st in STATES:
        for _, x in t3[t3.state == st].iterrows():
            put(ws, r, 1, NAMES[st]); put(ws, r, 2, x.level); put(ws, r, 3, int(x.total), fmt=CNT); put(ws, r, 4, int(x.done), fmt=CNT)
            put(ws, r, 5, int(x.live), fmt=CNT); put(ws, r, 6, int(x.dead), fmt=CNT); put(ws, r, 7, f"=D{r}/C{r}", fmt=PCT)
            if fc is not None and pd.notna(x.done_adj):
                put(ws, r, 8, int(x.done_adj), fmt=CNT); put(ws, r, 9, f"=H{r}/C{r}", fmt=PCT)
            r += 1
        r += 1
    if fc is not None:
        r += 1; put(ws, r, 1, f"City feeds that never record completions (done label on < {FEED_FLOOR:.0%} of the feed's raw permits), by state", bold=True); r += 1
        header(ws, r, ["State", "City feeds in scan", "Feeds not closing", "Raw permits in non-closing feeds", "Share of raw permits"]); r += 1
        for st in STATES:
            g = fc[fc.state == st]
            put(ws, r, 1, NAMES[st]); put(ws, r, 2, len(g), fmt=CNT); put(ws, r, 3, int((~g.closes).sum()), fmt=CNT)
            put(ws, r, 4, int(g.loc[~g.closes, "n"].sum()), fmt=CNT); put(ws, r, 5, f"=D{r}/{int(g.n.sum())}", fmt=PCT)
            r += 1
    widths(ws, first=14, rest=18); ws.column_dimensions["B"].width = 46

def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--data", required=True); ap.add_argument("--out", required=True); ap.add_argument("--scan")
    a = ap.parse_args(); con = duckdb.connect()
    pe = pd.read_stata(f"{a.data}/clean/permit_elevations.dta"); pe = pe[pe.state.isin(STATES)].copy()
    t1 = table1(con, a.data); ta, tb = table2(con, a.data, pe)
    fc = feed_closes(a.scan)
    t3 = table3(con, a.data, pe, fc)
    wb = Workbook(); wb.remove(wb.active)
    wb.calculation.fullCalcOnLoad = True   # no LibreOffice on this Mac: Excel computes the formulas on open
    sheet1(wb, t1); sheet2(wb, ta, tb); sheet3(wb, t3, fc)
    wb.save(a.out); print("wrote", a.out)
    pd.set_option("display.width", 220); print(t1.to_string(index=False)); print(ta.to_string(index=False)); print(t3.to_string(index=False))

if __name__ == "__main__":
    main()
